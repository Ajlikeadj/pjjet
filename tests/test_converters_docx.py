"""DOCX converter: order, hyperlinks, tables, footnotes, tracked changes and guards."""

from __future__ import annotations

import io
import json
import sys
import zipfile

import pytest

import pjjet

docx = pytest.importorskip("docx")


def build(builder) -> bytes:  # type: ignore[no-untyped-def]
    doc = docx.Document()
    builder(doc)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def convert(data: bytes, *, name: str | None = "a.docx", **options):  # type: ignore[no-untyped-def]
    options.setdefault("token_backend", "chars")
    return pjjet.convert(data, name=name, options=pjjet.Options(**options))


def codes(result) -> list[str]:  # type: ignore[no-untyped-def]
    return [w.code for w in result.warnings]


def with_part(data: bytes, part_name: str, content: bytes) -> bytes:
    """Rewrite a DOCX with one extra (unreferenced) package part."""
    source = zipfile.ZipFile(io.BytesIO(data))
    target = io.BytesIO()
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as archive:
        for info in source.infolist():
            archive.writestr(info, source.read(info.filename))
        archive.writestr(part_name, content)
    return target.getvalue()


# ── Structure ──────────────────────────────────────────────────────────────


def test_document_order_is_preserved():
    def builder(doc):  # type: ignore[no-untyped-def]
        doc.add_heading("Title", level=1)
        doc.add_paragraph("Before table.")
        table = doc.add_table(rows=3, cols=2)
        for row, (left, right) in enumerate([("Item", "Qty"), ("Widget", "3"), ("Gadget", "10")]):
            table.cell(row, 0).text = left
            table.cell(row, 1).text = right
        doc.add_heading("After", level=2)

    r = convert(build(builder))
    positions = [r.text.index(x) for x in ("Title", "Before table", "Widget", "After")]
    assert positions == sorted(positions)
    assert "# Title" in r.text and "## After" in r.text
    assert r.warnings == ()


def test_hyperlinks_are_rendered_as_markdown():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def builder(doc):  # type: ignore[no-untyped-def]
        paragraph = doc.add_paragraph("See ")
        rid = doc.part.relate_to(
            "https://example.com/report",
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            is_external=True,
        )
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), rid)
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "the report"
        run.append(text)
        link.append(run)
        paragraph._p.append(link)

    assert "See [the report](https://example.com/report)" in convert(build(builder)).text


def test_tabs_and_line_breaks_survive():
    def builder(doc):  # type: ignore[no-untyped-def]
        paragraph = doc.add_paragraph()
        paragraph.add_run("a\tb")
        paragraph.add_run().add_break()
        paragraph.add_run("c")

    assert "a\tb\nc" in convert(build(builder)).text


def test_empty_document():
    r = convert(build(lambda doc: None))
    assert r.text == "" and r.warnings == () and r.stats.converter == "docx"


def test_paragraph_style_heading_levels():
    def builder(doc):  # type: ignore[no-untyped-def]
        doc.add_paragraph("One", style="Heading 1")
        doc.add_paragraph("Six", style="Heading 6")

    assert convert(build(builder)).text == "# One\n\n###### Six\n"


# ── Tables ─────────────────────────────────────────────────────────────────


def test_single_row_table_is_data_not_header():
    def builder(doc):  # type: ignore[no-untyped-def]
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "a"
        table.cell(0, 1).text = "b"

    r = convert(build(builder), format="json-compact")
    block = json.loads(r.text)[0]
    assert "header" not in block
    assert block["rows"] == [["a", "b"]]
    assert "single-row-table" in codes(convert(build(builder)))


def test_empty_table_is_skipped():
    def builder(doc):  # type: ignore[no-untyped-def]
        doc.add_table(rows=0, cols=2)
        doc.add_paragraph("after")

    r = convert(build(builder))
    assert r.text == "after\n"


def test_merged_cells_are_reported():
    def builder(doc):  # type: ignore[no-untyped-def]
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "merged"
        table.cell(1, 0).text = "merged"
        table.cell(0, 0).merge(table.cell(1, 0))
        table.cell(0, 1).text = "x"
        table.cell(1, 1).text = "y"

    assert "merged-cells" in codes(convert(build(builder)))


# ── Footnotes and tracked changes ──────────────────────────────────────────

FOOTNOTES = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:t>sep</w:t></w:r></w:p></w:footnote>
  <w:footnote w:id="1"><w:p><w:r><w:t>First note</w:t></w:r></w:p></w:footnote>
  <w:footnote w:id="2"><w:p><w:r><w:t>Second note</w:t></w:r></w:p></w:footnote>
</w:footnotes>
"""


def test_footnotes_are_extracted():
    data = with_part(build(lambda doc: doc.add_paragraph("body")), "word/footnotes.xml", FOOTNOTES)
    r = convert(data)
    assert "footnote: 1: First note" in r.text
    assert "footnote: 2: Second note" in r.text
    assert "sep" not in r.text


def test_tracked_insertions_are_accepted_and_deletions_ignored():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def builder(doc):  # type: ignore[no-untyped-def]
        paragraph = doc.add_paragraph("Kept ")
        ins = OxmlElement("w:ins")
        ins.set(qn("w:id"), "1")
        ins.set(qn("w:author"), "tester")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "inserted"
        run.append(text)
        ins.append(run)
        paragraph._p.append(ins)

        deletion = OxmlElement("w:del")
        deletion.set(qn("w:id"), "2")
        deleted_run = OxmlElement("w:r")
        deleted_text = OxmlElement("w:delText")
        deleted_text.text = "deleted"
        deleted_run.append(deleted_text)
        deletion.append(deleted_run)
        paragraph._p.append(deletion)

    r = convert(build(builder))
    assert "Kept inserted" in r.text
    assert "deleted" not in r.text
    assert "tracked-changes" in codes(r)
    assert r.warnings[0].severity == "info"


# ── Corruption and guards ──────────────────────────────────────────────────


def test_non_zip_docx_is_rejected_without_crashing():
    r = convert(b"PK\x03\x04this is not really a zip file", name="broken.docx")
    assert r.text == ""
    assert codes(r) == ["docx-invalid"]
    assert r.warnings[0].severity == "skip"


def test_plain_text_forced_as_docx_is_rejected():
    r = convert(b"just some text", name="a.txt", type="docx")
    assert codes(r) == ["docx-invalid"]


def test_zip_bomb_ratio_is_refused():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"\x00" * (4 * 1024 * 1024))
    r = convert(buffer.getvalue(), name="bomb.docx")
    assert r.text == ""
    assert codes(r) == ["docx-invalid"]
    assert "zip bomb" in r.warnings[0].message


def test_missing_python_docx_is_a_warning(monkeypatch):
    monkeypatch.setitem(sys.modules, "docx", None)
    r = convert(b"PK\x03\x04", name="a.docx", type="docx")
    assert r.text == ""
    assert codes(r) == ["docx-unavailable"]
    assert "pip install" in r.warnings[0].message


def test_page_like_structure_does_not_invent_tracked_changes():
    # Regression: every run used to be counted as a tracked change.
    r = convert(build(lambda doc: doc.add_paragraph("a").add_run("b")))
    assert "tracked-changes" not in codes(r)
    assert r.text == "ab\n"


def test_unicode_and_xml_entities_are_preserved():
    r = convert(build(lambda doc: doc.add_paragraph("caf\u00e9 & <tag> \u65e5\u672c\u8a9e")))
    assert r.text == "caf\u00e9 & <tag> \u65e5\u672c\u8a9e\n"


def test_docx_metadata_is_reported():
    r = convert(build(lambda doc: doc.add_paragraph("x")), front_matter=True)
    assert r.text.startswith("---\nsource: a.docx\ntype: docx\n")
    assert "tokens:" in r.text and "token_backend:" in r.text
