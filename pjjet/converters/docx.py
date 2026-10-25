"""DOCX converter for pjjet.

Extracts headings, paragraphs (with inline hyperlinks as ``[text](url)``), tables,
footnotes and a summary of tracked changes. Tracked insertions are accepted; tracked
deletions are excluded.

Security: the container is inspected before python-docx parses it. Oversized archives and
absurd compression ratios (zip bombs) are refused, and ``word/footnotes.xml`` is parsed
with ``defusedxml`` when installed.
"""

from __future__ import annotations

import re
import zipfile
from typing import Any, BinaryIO

from ..ir import Block, ConversionWarning, Document, Heading, KeyValue, Meta, Paragraph
from ..ir import Table as IRTable
from ..options import Options

#: Refuse archives that expand past this size (a DOCX is a ZIP file).
MAX_UNCOMPRESSED_BYTES = 512 * 1024 * 1024
#: Refuse members with an implausible compression ratio.
MAX_COMPRESSION_RATIO = 200
_MIN_RATIO_FILE = 1_000_000

_HEADING_STYLES = {
    "heading 1": 1,
    "heading 2": 2,
    "heading 3": 3,
    "heading 4": 4,
    "heading 5": 5,
    "heading 6": 6,
    "title": 1,
    "subtitle": 2,
}
_SKIP_TAGS = {
    "pPr",
    "rPr",
    "del",
    "delText",
    "instrText",
    "proofErr",
    "bookmarkStart",
    "bookmarkEnd",
}
_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_MAX_FOOTNOTES = 50


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _clean(text: str) -> str:
    return text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")


def _zip_guard(stream: BinaryIO) -> str | None:
    """Return a refusal message when the container is malformed or dangerous."""
    stream.seek(0)
    try:
        if not zipfile.is_zipfile(stream):
            return "not a valid ZIP/DOCX container"
        stream.seek(0)
        with zipfile.ZipFile(stream) as archive:
            total = 0
            for info in archive.infolist():
                total += info.file_size
                if total > MAX_UNCOMPRESSED_BYTES:
                    return (
                        "archive expands beyond "
                        f"{MAX_UNCOMPRESSED_BYTES // (1024 * 1024)} MiB (zip bomb guard)"
                    )
                if (
                    info.file_size > _MIN_RATIO_FILE
                    and info.compress_size
                    and info.file_size / info.compress_size > MAX_COMPRESSION_RATIO
                ):
                    return "archive has an implausible compression ratio (zip bomb guard)"
    except (zipfile.BadZipFile, OSError) as exc:
        return f"not a valid ZIP/DOCX container: {exc}"
    finally:
        stream.seek(0)
    return None


def _run_text(element: Any) -> str:
    """Text of a ``w:r`` including tabs/breaks; deleted text is excluded."""
    parts: list[str] = []
    for node in element.iter():
        name = _local(node.tag)
        if name == "t":
            parts.append(node.text or "")
        elif name == "tab":
            parts.append("\t")
        elif name in ("br", "cr"):
            parts.append("\n")
    return "".join(parts)


def _hyperlink_target(paragraph: Any, element: Any) -> str | None:
    from docx.oxml.ns import qn

    rid = element.get(qn("r:id"))
    if not rid:
        return None
    try:
        return str(paragraph.part.rels[rid].target_ref)
    except Exception:  # missing/broken relationship: keep the text, drop the link
        return None


def _paragraph_text(paragraph: Any) -> str:
    pieces: list[str] = []

    def walk(node: Any) -> None:
        for child in node:
            name = _local(child.tag)
            if name in _SKIP_TAGS:
                continue
            if name == "r":
                pieces.append(_run_text(child))
            elif name == "hyperlink":
                text = "".join(_run_text(run) for run in child if _local(run.tag) == "r")
                target = _hyperlink_target(paragraph, child)
                pieces.append(f"[{text}]({target})" if text and target else text)
            else:
                walk(child)

    walk(paragraph._p)
    return _clean("".join(pieces))


def _classify_heading(style: str) -> int:
    style = style.strip().lower()
    if style in _HEADING_STYLES:
        return _HEADING_STYLES[style]
    match = re.fullmatch(r"heading\s+(\d+)", style)
    if match:
        return min(max(int(match.group(1)), 1), 6)
    return 0


def _heading_level(paragraph: Any) -> int:
    style = getattr(getattr(paragraph, "style", None), "name", "") or ""
    level = _classify_heading(style)
    if level:
        return level
    ppr = paragraph._p.pPr
    if ppr is not None:
        outline = ppr.find(f"{_W_NS}outlineLvl")
        if outline is not None:
            try:
                return min(max(int(outline.get(f"{_W_NS}val", "")) + 1, 1), 6)
            except (TypeError, ValueError):
                return 0
    return 0


def _has_merged_cells(table: Any) -> bool:
    from docx.oxml.ns import qn

    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._element.find(qn("w:tcPr"))
            if tc_pr is None:
                continue
            if tc_pr.find(qn("w:vMerge")) is not None or tc_pr.find(qn("w:gridSpan")) is not None:
                return True
    return False


def _fromstring(data: bytes) -> Any:
    try:
        from defusedxml.ElementTree import fromstring
    except ImportError:  # pragma: no cover - defusedxml is a declared extra
        from xml.etree.ElementTree import fromstring
    return fromstring(data)


def _footnotes(stream: BinaryIO) -> tuple[str, ...]:
    """Read ``word/footnotes.xml``; python-docx does not expose footnotes."""
    from docx.oxml.ns import qn

    stream.seek(0)
    try:
        with zipfile.ZipFile(stream) as archive:
            try:
                data = archive.read("word/footnotes.xml")
            except KeyError:
                return ()
    except (zipfile.BadZipFile, OSError):
        return ()
    finally:
        stream.seek(0)

    try:
        root = _fromstring(data)
    except Exception:
        return ()
    notes: list[str] = []
    for note in root.iter(qn("w:footnote")):
        if note.get(qn("w:type")) in ("separator", "continuationSeparator", "continuationNotice"):
            continue
        text = "".join(node.text or "" for node in note.iter(qn("w:t"))).strip()
        if text:
            notes.append(text)
    return tuple(notes[:_MAX_FOOTNOTES])


class DocxConverter:
    """Converts ``.docx`` files into clean IR blocks."""

    name = "docx"
    extensions = (".docx",)
    needs_seekable = True

    def detect(self, head: bytes, filename: str | None) -> float:
        if filename and filename.lower().endswith(self.extensions):
            return 1.0
        if head and head[:4] == b"PK\x03\x04":
            return 0.5
        return 0.0

    def convert(self, stream: BinaryIO, opts: Options, filename: str | None) -> Document:
        warnings: list[ConversionWarning] = []
        meta = Meta(source=filename, type=self.name)

        try:
            from docx import Document as DocxDocument
            from docx.table import Table as DocxTable
            from docx.text.paragraph import Paragraph as DocxParagraph
        except ImportError as exc:
            warnings.append(
                ConversionWarning(
                    "docx-unavailable",
                    f"python-docx is not installed ({exc}); install with "
                    "`pip install 'pjjet[docx]'`.",
                    "skip",
                )
            )
            return Document([], meta, warnings)

        refusal = _zip_guard(stream)
        if refusal is not None:
            warnings.append(
                ConversionWarning("docx-invalid", f"Cannot read DOCX: {refusal}", "skip")
            )
            return Document([], meta, warnings)

        stream.seek(0)
        try:
            doc = DocxDocument(stream)
        except Exception as exc:
            warnings.append(
                ConversionWarning(
                    "docx-error",
                    f"python-docx could not open the file: {type(exc).__name__}: {exc}",
                    "skip",
                )
            )
            return Document([], meta, warnings)

        blocks: list[Block] = []
        try:
            items = list(doc.iter_inner_content())
        except AttributeError:  # python-docx < 1.1
            items = [*doc.paragraphs, *doc.tables]

        for item in items:
            if isinstance(item, DocxParagraph):
                text = _paragraph_text(item)
                if not text.strip():
                    continue
                level = _heading_level(item)
                blocks.append(Heading(level=level, text=text) if level else Paragraph(text=text))
            elif isinstance(item, DocxTable):
                block = self._table(item, opts, warnings)
                if block is not None:
                    blocks.append(block)

        for index, note in enumerate(_footnotes(stream), 1):
            blocks.append(KeyValue((("footnote", f"{index}: {note}"),)))

        changes = self._tracked_changes(doc)
        if changes:
            warnings.append(
                ConversionWarning(
                    "tracked-changes",
                    f"{changes} tracked change(s): insertions accepted, deletions ignored.",
                    "info",
                )
            )

        return Document(blocks, meta, warnings)

    def _table(
        self, table: Any, opts: Options, warnings: list[ConversionWarning]
    ) -> IRTable | None:
        try:
            matrix = [[_clean(cell.text) for cell in row.cells] for row in table.rows]
        except Exception as exc:
            warnings.append(
                ConversionWarning(
                    "table-error",
                    f"Skipped a table that could not be read: {type(exc).__name__}: {exc}",
                    "skip",
                )
            )
            return None

        matrix = [row for row in matrix if any(cell.strip() for cell in row)]
        if not matrix:
            return None

        width = max(len(row) for row in matrix)
        matrix = [row + [""] * (width - len(row)) for row in matrix]

        header: tuple[str, ...] | None = None
        data = matrix
        first = [cell.strip() for cell in matrix[0]]
        if (
            len(matrix) > 1
            and all(first)
            and len(set(first)) == len(first)
            and not any(value.replace(".", "", 1).isdigit() for value in first)
        ):
            header = tuple(matrix[0])
            data = matrix[1:]

        if _has_merged_cells(table):
            warnings.append(
                ConversionWarning(
                    "merged-cells",
                    "Table contains merged cells; the same value may repeat across cells.",
                    "info",
                )
            )
        if len(matrix) == 1:
            warnings.append(
                ConversionWarning(
                    "single-row-table",
                    "Table has a single row; it was treated as data, not a header.",
                    "info",
                )
            )
        return IRTable(
            header=header,
            rows=tuple(tuple(row) for row in data),
            style=opts.table_style,
        )

    def _tracked_changes(self, doc: Any) -> int:
        from docx.oxml.ns import qn

        body = doc.element.body
        return sum(1 for _ in body.iter(qn("w:ins"))) + sum(1 for _ in body.iter(qn("w:del")))
